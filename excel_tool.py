#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Excel 批量处理大师 v1.0
功能：合并、拆分、去重、筛选、统计、格式化
作者：AI Assistant
"""

import os
import sys
import json
import threading
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext
from datetime import datetime

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'openpyxl'])
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

try:
    import pandas as pd
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'pandas'])
    import pandas as pd


class ExcelTool:
    """Excel 批量处理核心引擎"""

    def __init__(self, log_callback=None):
        self.log = log_callback or print

    def merge_files(self, file_list, output_path, ignore_header=False):
        """合并多个 Excel 文件"""
        all_data = []
        for f in file_list:
            try:
                df = pd.read_excel(f)
                df['_来源文件'] = os.path.basename(f)
                all_data.append(df)
                self.log(f"  ✓ 读取: {os.path.basename(f)} ({len(df)}行)")
            except Exception as e:
                self.log(f"  ✗ 失败: {os.path.basename(f)} - {e}")

        if not all_data:
            self.log("错误：没有成功读取任何文件")
            return False

        result = pd.concat(all_data, ignore_index=True)
        if ignore_header:
            # 保留第一个文件的列名，后续文件去掉表头行（已在read_excel处理）
            pass

        result.to_excel(output_path, index=False, engine='openpyxl')
        self.log(f"\n合并完成！共 {len(result)} 行，保存到: {output_path}")
        return True

    def split_by_column(self, file_path, column, output_dir):
        """按某列的值拆分成多个文件"""
        df = pd.read_excel(file_path)
        if column not in df.columns:
            self.log(f"错误：找不到列 '{column}'")
            self.log(f"可用列: {', '.join(df.columns)}")
            return False

        os.makedirs(output_dir, exist_ok=True)
        groups = df.groupby(column)
        count = 0
        for name, group in groups:
            safe_name = str(name).replace('/', '_').replace('\\', '_')[:50]
            out_file = os.path.join(output_dir, f"{safe_name}.xlsx")
            group.drop(columns=[column]).to_excel(out_file, index=False, engine='openpyxl')
            count += 1
            self.log(f"  ✓ {safe_name}.xlsx ({len(group)}行)")

        self.log(f"\n拆分完成！共生成 {count} 个文件")
        return True

    def split_by_rows(self, file_path, rows_per_file, output_dir):
        """按行数拆分"""
        df = pd.read_excel(file_path)
        os.makedirs(output_dir, exist_ok=True)
        count = 0
        for i in range(0, len(df), rows_per_file):
            chunk = df.iloc[i:i + rows_per_file]
            out_file = os.path.join(output_dir, f"part_{count + 1}.xlsx")
            chunk.to_excel(out_file, index=False, engine='openpyxl')
            count += 1
            self.log(f"  ✓ part_{count}.xlsx ({len(chunk)}行)")

        self.log(f"\n拆分完成！共生成 {count} 个文件")
        return True

    def remove_duplicates(self, file_path, columns=None, output_path=None):
        """去重"""
        df = pd.read_excel(file_path)
        original = len(df)

        if columns:
            cols = [c for c in columns if c in df.columns]
            if not cols:
                self.log(f"错误：指定的列不存在")
                return False
            df = df.drop_duplicates(subset=cols)
        else:
            df = df.drop_duplicates()

        removed = original - len(df)
        if output_path is None:
            output_path = file_path.replace('.xlsx', '_去重后.xlsx')

        df.to_excel(output_path, index=False, engine='openpyxl')
        self.log(f"去重完成！原始 {original} 行，去除 {removed} 行，剩余 {len(df)} 行")
        return True

    def filter_rows(self, file_path, column, condition, value, output_path=None):
        """筛选数据"""
        df = pd.read_excel(file_path)
        if column not in df.columns:
            self.log(f"错误：找不到列 '{column}'")
            return False

        original = len(df)
        col = df[column]

        if condition == '包含':
            mask = col.astype(str).str.contains(str(value), na=False)
        elif condition == '不包含':
            mask = ~col.astype(str).str.contains(str(value), na=False)
        elif condition == '等于':
            mask = col.astype(str) == str(value)
        elif condition == '不等于':
            mask = col.astype(str) != str(value)
        elif condition == '大于':
            mask = pd.to_numeric(col, errors='coerce') > float(value)
        elif condition == '小于':
            mask = pd.to_numeric(col, errors='coerce') < float(value)
        elif condition == '为空':
            mask = col.isna()
        elif condition == '不为空':
            mask = col.notna()
        else:
            self.log(f"错误：不支持的条件 '{condition}'")
            return False

        result = df[mask]
        if output_path is None:
            output_path = file_path.replace('.xlsx', '_筛选结果.xlsx')

        result.to_excel(output_path, index=False, engine='openpyxl')
        self.log(f"筛选完成！原始 {original} 行，符合条件 {len(result)} 行")
        return True

    def statistics(self, file_path, group_column=None, value_column=None):
        """数据统计"""
        df = pd.read_excel(file_path)
        self.log(f"总行数: {len(df)}")
        self.log(f"总列数: {len(df.columns)}")
        self.log(f"列名: {', '.join(df.columns)}")
        self.log(f"\n数据类型:")
        for col in df.columns:
            self.log(f"  {col}: {df[col].dtype}")

        # 数值列统计
        numeric_cols = df.select_dtypes(include=['number']).columns
        if len(numeric_cols) > 0:
            self.log(f"\n数值列统计:")
            for col in numeric_cols:
                self.log(f"  {col}: 总计={df[col].sum():.2f}, 平均={df[col].mean():.2f}, "
                         f"最大={df[col].max():.2f}, 最小={df[col].min():.2f}")

        # 空值统计
        null_counts = df.isnull().sum()
        if null_counts.any():
            self.log(f"\n空值统计:")
            for col in df.columns:
                if null_counts[col] > 0:
                    self.log(f"  {col}: {null_counts[col]} 个空值")

        # 按列分组统计
        if group_column and group_column in df.columns:
            self.log(f"\n按 '{group_column}' 分组统计:")
            for name, group in df.groupby(group_column):
                self.log(f"  {name}: {len(group)} 行")

        return True

    def format_beautify(self, file_path, output_path=None):
        """美化格式"""
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active

        # 表头样式
        header_font = Font(bold=True, size=12, color='FFFFFF')
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_align = Alignment(horizontal='center', vertical='center')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # 设置表头
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

        # 设置数据行
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = thin_border
                cell.alignment = Alignment(vertical='center')
                # 隔行变色
                if cell.row % 2 == 0:
                    cell.fill = PatternFill(start_color='D9E2F3', end_color='D9E2F3', fill_type='solid')

        # 自动列宽
        for col in ws.columns:
            max_length = 0
            col_letter = col[0].column_letter
            for cell in col:
                try:
                    cell_len = len(str(cell.value or ''))
                    # 中文字符算2个宽度
                    cjk = sum(1 for c in str(cell.value or '') if '\u4e00' <= c <= '\u9fff')
                    cell_len += cjk
                    max_length = max(max_length, cell_len)
                except:
                    pass
            ws.column_dimensions[col_letter].width = min(max_length + 4, 50)

        # 冻结首行
        ws.freeze_panes = 'A2'

        if output_path is None:
            output_path = file_path.replace('.xlsx', '_美化版.xlsx')

        wb.save(output_path)
        self.log(f"美化完成！保存到: {output_path}")
        return True

    def vlookup(self, file1, file2, key_column, lookup_columns, output_path=None):
        """VLOOKUP 类似功能 - 跨表匹配"""
        df1 = pd.read_excel(file1)
        df2 = pd.read_excel(file2)

        if key_column not in df1.columns:
            self.log(f"错误：主表找不到列 '{key_column}'")
            return False
        if key_column not in df2.columns:
            self.log(f"错误：查找表找不到列 '{key_column}'")
            return False

        valid_cols = [c for c in lookup_columns if c in df2.columns]
        if not valid_cols:
            self.log(f"错误：查找表中没有指定的列")
            return False

        lookup_df = df2[[key_column] + valid_cols].drop_duplicates(subset=[key_column])
        result = df1.merge(lookup_df, on=key_column, how='left')

        if output_path is None:
            output_path = file1.replace('.xlsx', '_匹配结果.xlsx')

        result.to_excel(output_path, index=False, engine='openpyxl')
        self.log(f"匹配完成！主表 {len(df1)} 行，匹配后保存到: {output_path}")
        return True


class ExcelToolGUI:
    """图形界面"""

    def __init__(self):
        self.root = tk.Tk()
        self.root.title("Excel 批量处理大师 v1.0")
        self.root.geometry("800x650")
        self.root.resizable(True, True)

        # 设置图标和样式
        style = ttk.Style()
        style.theme_use('clam')

        self.tool = ExcelTool(log_callback=self.log)
        self._build_ui()

    def _build_ui(self):
        """构建界面"""
        # 主框架
        main = ttk.Frame(self.root, padding=10)
        main.pack(fill=tk.BOTH, expand=True)

        # 功能选择
        func_frame = ttk.LabelFrame(main, text="选择功能", padding=10)
        func_frame.pack(fill=tk.X, pady=(0, 10))

        self.func_var = tk.StringVar(value='merge')
        functions = [
            ('merge', '合并文件', '将多个Excel合并为一个'),
            ('split_col', '按列拆分', '按某列的值拆成多个文件'),
            ('split_row', '按行拆分', '按固定行数拆分'),
            ('dedup', '去重', '删除重复行'),
            ('filter', '筛选数据', '按条件筛选行'),
            ('stats', '数据统计', '查看数据概览'),
            ('beautify', '美化格式', '自动美化Excel样式'),
            ('vlookup', '跨表匹配', '类似VLOOKUP功能'),
        ]

        for i, (val, label, desc) in enumerate(functions):
            r, c = divmod(i, 4)
            frame = ttk.Frame(func_frame)
            frame.grid(row=r, column=c, padx=5, pady=2, sticky='w')
            ttk.Radiobutton(frame, text=label, variable=self.func_var,
                            value=val, command=self._update_options).pack(side=tk.LEFT)
            ttk.Label(frame, text=desc, foreground='gray').pack(side=tk.LEFT, padx=(5, 0))

        # 选项区域
        self.opt_frame = ttk.LabelFrame(main, text="参数设置", padding=10)
        self.opt_frame.pack(fill=tk.X, pady=(0, 10))

        # 文件选择区
        file_frame = ttk.LabelFrame(main, text="文件选择", padding=10)
        file_frame.pack(fill=tk.X, pady=(0, 10))

        self.file_list = []

        btn_row = ttk.Frame(file_frame)
        btn_row.pack(fill=tk.X)

        ttk.Button(btn_row, text="添加文件", command=self._add_files).pack(side=tk.LEFT, padx=2)
        ttk.Button(btn_row, text="添加文件夹", command=self._add_folder).pack(side=tk.LEFT, padx=2)
        ttk.Button(btn_row, text="清空", command=self._clear_files).pack(side=tk.LEFT, padx=2)
        ttk.Button(btn_row, text="选择输出位置", command=self._select_output).pack(side=tk.RIGHT, padx=2)

        self.file_label = ttk.Label(file_frame, text="未选择文件", foreground='gray')
        self.file_label.pack(anchor='w', pady=(5, 0))
        self.output_path = ''

        # 执行按钮
        btn_frame = ttk.Frame(main)
        btn_frame.pack(fill=tk.X, pady=(0, 10))
        ttk.Button(btn_frame, text="▶ 开始执行", command=self._execute).pack(side=tk.LEFT)
        ttk.Button(btn_frame, text="清空日志", command=self._clear_log).pack(side=tk.RIGHT)

        # 日志区
        log_frame = ttk.LabelFrame(main, text="执行日志", padding=5)
        log_frame.pack(fill=tk.BOTH, expand=True)
        self.log_text = scrolledtext.ScrolledText(log_frame, height=10, font=('Consolas', 9))
        self.log_text.pack(fill=tk.BOTH, expand=True)

        self._update_options()

    def _update_options(self):
        """根据选择的功能更新参数选项"""
        for w in self.opt_frame.winfo_children():
            w.destroy()

        func = self.func_var.get()

        if func == 'merge':
            self.ignore_header_var = tk.BooleanVar()
            ttk.Checkbutton(self.opt_frame, text="忽略重复表头", variable=self.ignore_header_var).pack(anchor='w')

        elif func == 'split_col':
            ttk.Label(self.opt_frame, text="拆分依据列名:").pack(anchor='w')
            self.split_col_var = tk.StringVar()
            ttk.Entry(self.opt_frame, textvariable=self.split_col_var, width=30).pack(anchor='w')

        elif func == 'split_row':
            ttk.Label(self.opt_frame, text="每个文件行数:").pack(anchor='w')
            self.split_rows_var = tk.StringVar(value='1000')
            ttk.Entry(self.opt_frame, textvariable=self.split_rows_var, width=10).pack(anchor='w')

        elif func == 'dedup':
            ttk.Label(self.opt_frame, text="去重依据列名（留空=全列去重，多列用逗号分隔）:").pack(anchor='w')
            self.dedup_cols_var = tk.StringVar()
            ttk.Entry(self.opt_frame, textvariable=self.dedup_cols_var, width=50).pack(anchor='w')

        elif func == 'filter':
            row1 = ttk.Frame(self.opt_frame)
            row1.pack(fill=tk.X, pady=2)
            ttk.Label(row1, text="列名:").pack(side=tk.LEFT)
            self.filter_col_var = tk.StringVar()
            ttk.Entry(row1, textvariable=self.filter_col_var, width=20).pack(side=tk.LEFT, padx=5)

            ttk.Label(row1, text="条件:").pack(side=tk.LEFT, padx=(10, 0))
            self.filter_cond_var = tk.StringVar(value='包含')
            ttk.Combobox(row1, textvariable=self.filter_cond_var, width=8,
                         values=['包含', '不包含', '等于', '不等于', '大于', '小于', '为空', '不为空'],
                         state='readonly').pack(side=tk.LEFT, padx=5)

            ttk.Label(row1, text="值:").pack(side=tk.LEFT, padx=(10, 0))
            self.filter_val_var = tk.StringVar()
            ttk.Entry(row1, textvariable=self.filter_val_var, width=20).pack(side=tk.LEFT, padx=5)

        elif func == 'vlookup':
            ttk.Label(self.opt_frame, text="匹配键列名（两个表共有的列）:").pack(anchor='w')
            self.vlookup_key_var = tk.StringVar()
            ttk.Entry(self.opt_frame, textvariable=self.vlookup_key_var, width=30).pack(anchor='w')
            ttk.Label(self.opt_frame, text="要查找的列名（逗号分隔）:").pack(anchor='w', pady=(5, 0))
            self.vlookup_cols_var = tk.StringVar()
            ttk.Entry(self.opt_frame, textvariable=self.vlookup_cols_var, width=50).pack(anchor='w')
            ttk.Label(self.opt_frame, text="注意：第一个文件=主表，第二个文件=查找表",
                      foreground='red').pack(anchor='w', pady=(5, 0))

    def _add_files(self):
        files = filedialog.askopenfilenames(
            title="选择Excel文件",
            filetypes=[("Excel文件", "*.xlsx *.xls"), ("所有文件", "*.*")]
        )
        self.file_list.extend(files)
        self._update_file_label()

    def _add_folder(self):
        folder = filedialog.askdirectory(title="选择文件夹")
        if folder:
            for f in os.listdir(folder):
                if f.endswith(('.xlsx', '.xls')) and not f.startswith('~'):
                    self.file_list.append(os.path.join(folder, f))
            self._update_file_label()

    def _clear_files(self):
        self.file_list = []
        self._update_file_label()

    def _update_file_label(self):
        if self.file_list:
            self.file_label.config(text=f"已选择 {len(self.file_list)} 个文件", foreground='blue')
        else:
            self.file_label.config(text="未选择文件", foreground='gray')

    def _select_output(self):
        func = self.func_var.get()
        if func in ('merge', 'dedup', 'filter', 'beautify', 'vlookup'):
            path = filedialog.asksaveasfilename(
                defaultextension='.xlsx',
                filetypes=[("Excel文件", "*.xlsx")],
                title="保存输出文件"
            )
        else:
            path = filedialog.askdirectory(title="选择输出文件夹")
        if path:
            self.output_path = path

    def _execute(self):
        """执行功能"""
        func = self.func_var.get()

        if not self.file_list:
            messagebox.showwarning("提示", "请先选择文件！")
            return

        # 在线程中执行
        threading.Thread(target=self._run_task, args=(func,), daemon=True).start()

    def _run_task(self, func):
        """在线程中运行任务"""
        try:
            if func == 'merge':
                out = self.output_path or os.path.join(os.path.dirname(self.file_list[0]), '合并结果.xlsx')
                self.tool.merge_files(self.file_list, out, self.ignore_header_var.get())

            elif func == 'split_col':
                col = self.split_col_var.get().strip()
                if not col:
                    self.log("请输入拆分依据的列名！")
                    return
                out_dir = self.output_path or os.path.join(os.path.dirname(self.file_list[0]), '拆分结果')
                self.tool.split_by_column(self.file_list[0], col, out_dir)

            elif func == 'split_row':
                rows = int(self.split_rows_var.get())
                out_dir = self.output_path or os.path.join(os.path.dirname(self.file_list[0]), '拆分结果')
                self.tool.split_by_rows(self.file_list[0], rows, out_dir)

            elif func == 'dedup':
                cols_str = self.dedup_cols_var.get().strip()
                cols = [c.strip() for c in cols_str.split(',') if c.strip()] if cols_str else None
                out = self.output_path or None
                self.tool.remove_duplicates(self.file_list[0], cols, out)

            elif func == 'filter':
                col = self.filter_col_var.get().strip()
                cond = self.filter_cond_var.get()
                val = self.filter_val_var.get().strip()
                out = self.output_path or None
                self.tool.filter_rows(self.file_list[0], col, cond, val, out)

            elif func == 'stats':
                self.tool.statistics(self.file_list[0])

            elif func == 'beautify':
                out = self.output_path or None
                self.tool.format_beautify(self.file_list[0], out)

            elif func == 'vlookup':
                key = self.vlookup_key_var.get().strip()
                cols = [c.strip() for c in self.vlookup_cols_var.get().split(',') if c.strip()]
                if not key or not cols:
                    self.log("请填写匹配键和查找列！")
                    return
                out = self.output_path or None
                self.tool.vlookup(self.file_list[0], self.file_list[1], key, cols, out)

        except Exception as e:
            self.log(f"\n✗ 执行出错: {e}")

    def log(self, msg):
        """写日志"""
        def _write():
            self.log_text.insert(tk.END, msg + '\n')
            self.log_text.see(tk.END)
        self.root.after(0, _write)

    def _clear_log(self):
        self.log_text.delete(1.0, tk.END)

    def run(self):
        self.log("Excel 批量处理大师 v1.0 就绪")
        self.log("请选择功能 → 添加文件 → 开始执行\n")
        self.root.mainloop()


if __name__ == '__main__':
    app = ExcelToolGUI()
    app.run()
